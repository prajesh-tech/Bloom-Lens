from typing import Dict, Any, List
import docx
from app.core.logging import logger
from app.utils.text_cleaner import clean_text


class DOCXService:
    """Service for extracting text from Microsoft Word DOCX question papers."""

    @staticmethod
    def extract_text_from_docx(file_path: str) -> Dict[str, Any]:
        """
        Extracts structured text from DOCX paragraphs and tables.
        """
        paragraphs_text: List[str] = []

        try:
            doc = docx.Document(file_path)

            # 1. Process regular paragraphs
            for p in doc.paragraphs:
                txt = clean_text(p.text)
                if txt:
                    paragraphs_text.append(txt)

            # 2. Process text inside table cells
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
                    if row_texts:
                        paragraphs_text.append(" | ".join(row_texts))

            full_text = "\n\n".join(paragraphs_text)
            logger.info(f"Extracted {len(full_text)} characters from DOCX: {file_path}")

            return {
                "pages": [{"page_number": 1, "text": full_text}],
                "full_text": full_text,
                "page_count": 1,
                "extraction_method": "docx_python_docx"
            }
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {e}")
            raise RuntimeError(f"Failed to process DOCX file: {str(e)}")
